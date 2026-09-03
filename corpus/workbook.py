import hashlib
from io import BytesIO
import json
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .mapping import MappingFixture


IMPORTER_VERSION = "m2-v1"


@dataclass(frozen=True)
class CellSnapshot:
    coordinate: str
    column_name: str
    data_type: str
    raw_value: str | None
    displayed_value: str | None
    formula: str | None


@dataclass
class SourceRowPlan:
    sheet_name: str
    row_number: int
    source_role: str
    classification: str
    interpretation_source: str
    structural_kind: str = ""
    cells: list[CellSnapshot] = field(default_factory=list)
    mapping_context: dict = field(default_factory=dict)
    domain_slug: str | None = None
    question_stable_id: str | None = None
    source_fingerprint: str = ""


@dataclass(frozen=True)
class DomainPlan:
    name: str
    slug: str
    code: str
    source_sheet: str
    start_row: int
    end_row: int
    interpretation_source: str
    meaningful_source_rows: int
    merged_source_rows: int
    question_count: int
    first_id: str
    last_id: str


@dataclass(frozen=True)
class QuestionPlan:
    stable_id: str
    question_text: str
    domain_slug: str
    tags: tuple[str, ...]
    source_role: str
    source_rows: tuple[int, ...]
    sheet_name: str


@dataclass(frozen=True)
class LegacyObservationPlan:
    sheet_name: str
    row_number: int
    question_stable_id: str | None
    source_question_text: str
    observed_answer_text: str | None
    legacy_expectation: str | None
    acceptable_source_value: str | None
    legacy_judgment: str | None
    comments: str | None
    secondary_comments: str | None
    experiment_metadata: dict


@dataclass(frozen=True)
class WarningPlan:
    sheet_name: str
    source_rows: tuple[int, ...]
    code: str
    explanation: str


@dataclass
class WorkbookImportPlan:
    source_path: Path
    source_filename: str
    source_sha256: str
    source_size_bytes: int
    source_metadata: dict
    inventory: dict
    mapping: MappingFixture
    domains: list[DomainPlan]
    questions: list[QuestionPlan]
    source_rows: list[SourceRowPlan]
    observations: list[LegacyObservationPlan]
    warnings: list[WarningPlan]
    tags: tuple[str, ...]


def _value_as_text(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


def _row_fingerprint(row):
    payload = {
        "sheet": row.sheet_name,
        "row": row.row_number,
        "cells": [
            {
                "coordinate": cell.coordinate,
                "column_name": cell.column_name,
                "data_type": cell.data_type,
                "raw_value": cell.raw_value,
                "displayed_value": cell.displayed_value,
                "formula": cell.formula,
            }
            for cell in row.cells
        ],
    }
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()


def _core_property_metadata(source_bytes):
    namespaces = {
        "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
        "dc": "http://purl.org/dc/elements/1.1/",
        "dcterms": "http://purl.org/dc/terms/",
    }
    with ZipFile(BytesIO(source_bytes), "r") as archive:
        root = ElementTree.fromstring(archive.read("docProps/core.xml"))

    def value(path):
        element = root.find(path, namespaces)
        return element.text if element is not None else None

    return {
        "creator": value("dc:creator"),
        "last_modified_by": value("cp:lastModifiedBy"),
        "created": value("dcterms:created"),
        "modified": value("dcterms:modified"),
    }


def _nonblank(value):
    return value is not None and (not isinstance(value, str) or value != "")


def _cell_snapshots(source_sheet, displayed_sheet, row_number, headers):
    snapshots = []
    for column_number, column_name in enumerate(headers, start=1):
        source_cell = source_sheet.cell(row=row_number, column=column_number)
        displayed_cell = displayed_sheet.cell(row=row_number, column=column_number)
        formula = _value_as_text(source_cell.value) if source_cell.data_type == "f" else None
        snapshots.append(
            CellSnapshot(
                coordinate=f"{get_column_letter(column_number)}{row_number}",
                column_name=column_name,
                data_type=source_cell.data_type or "",
                raw_value=_value_as_text(source_cell.value),
                displayed_value=_value_as_text(displayed_cell.value),
                formula=formula,
            )
        )
    return snapshots


def _row_values(sheet, row_number, width):
    return [sheet.cell(row=row_number, column=column).value for column in range(1, width + 1)]


def _validate_workbook_structure(source_book, displayed_book, mapping):
    expected_sheets = mapping.data["recognized_sheets"]
    if source_book.sheetnames != expected_sheets or displayed_book.sheetnames != expected_sheets:
        raise ValidationError(
            f"Workbook sheets differ from mapping: found {source_book.sheetnames!r}."
        )
    for sheet_name in expected_sheets:
        sheet = source_book[sheet_name]
        if sheet.sheet_state != "visible":
            raise ValidationError(f"Sheet {sheet_name!r} must be visible.")
        if sheet.merged_cells.ranges:
            raise ValidationError(f"Sheet {sheet_name!r} unexpectedly contains merged cells.")
        hidden_rows = [
            number
            for number, dimension in sheet.row_dimensions.items()
            if dimension.hidden
        ]
        hidden_columns = [
            name
            for name, dimension in sheet.column_dimensions.items()
            if dimension.hidden
        ]
        if hidden_rows or hidden_columns:
            raise ValidationError(
                f"Sheet {sheet_name!r} unexpectedly contains hidden rows/columns."
            )
        rules = mapping.data["sheets"][sheet_name]
        headers = rules.get("headers")
        if headers:
            actual = _row_values(sheet, 1, len(headers))
            if actual != headers:
                raise ValidationError(
                    f"Sheet {sheet_name!r} headers differ: expected {headers!r}, found {actual!r}."
                )
        if rules["role"] != "IGNORED" and sheet_name != "target questions":
            data_start, data_end = rules.get("data_rows", [2, sheet.max_row])
            width = len(headers) or 1
            blank_data_rows = [
                row_number
                for row_number in range(data_start, data_end + 1)
                if not any(
                    _nonblank(value)
                    for value in _row_values(sheet, row_number, width)
                )
            ]
            if blank_data_rows:
                raise ValidationError(
                    f"Sheet {sheet_name!r} contains unexpected blank mapped rows: "
                    f"{blank_data_rows}."
                )
    target = source_book["target questions"]
    if target.max_column != 1:
        raise ValidationError("target questions must contain only column A.")
    target_rules = mapping.data["sheets"]["target questions"]
    if _nonblank(target["A1"].value):
        raise ValidationError("target questions row 1 must remain unused.")
    for row_number in target_rules["separator_rows"]:
        if _nonblank(target.cell(row=row_number, column=1).value):
            raise ValidationError(f"Expected blank target separator at row {row_number}.")
    configured_rows = set()
    for domain in target_rules["domains"]:
        configured_rows.update(range(domain["start_row"], domain["end_row"] + 1))
    source_nonblank = {
        row_number
        for row_number in range(2, target.max_row + 1)
        if _nonblank(target.cell(row=row_number, column=1).value)
    }
    if source_nonblank != configured_rows:
        missing = sorted(configured_rows - source_nonblank)
        unexpected = sorted(source_nonblank - configured_rows)
        raise ValidationError(
            "Target row mapping differs from source; "
            f"blank configured={missing}, unexpected={unexpected}."
        )
    merge = target_rules["merge_rules"][0]
    joined = merge["join_with"].join(
        str(target.cell(row=row_number, column=1).value).strip()
        for row_number in merge["source_rows"]
    )
    if joined != merge["canonical_text"]:
        raise ValidationError(
            f"Rows 155-156 no longer match the approved canonical text: {joined!r}."
        )
    kb_rules = mapping.data["sheets"]["knowledge base - rag questions"]
    kb_sheet = source_book["knowledge base - rag questions"]
    kb_start, kb_end = kb_rules["data_rows"]
    kb_questions = {
        kb_sheet.cell(row=row_number, column=1).value
        for row_number in range(kb_start, kb_end + 1)
    }
    if len(kb_questions) != 1:
        raise ValidationError(
            "KB/RAG mapped rows no longer contain one exact repeated question."
        )


def _build_inventory(book, mapping):
    inventory = {}
    for sheet_name in mapping.data["recognized_sheets"]:
        sheet = book[sheet_name]
        rules = mapping.data["sheets"][sheet_name]
        headers = rules.get("headers", [])
        data_start, data_end = rules.get("data_rows", [2, sheet.max_row])
        width = len(headers) or 1
        meaningful = sum(
            any(_nonblank(value) for value in _row_values(sheet, row_number, width))
            for row_number in range(data_start, data_end + 1)
        )
        formulas = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas.append(cell.coordinate)
        inventory[sheet_name] = {
            "dimension": sheet.calculate_dimension(),
            "physical_rows": sheet.max_row,
            "physical_columns": sheet.max_column,
            "meaningful_rows": meaningful,
            "header_rows": 1 if headers else 0,
            "blank_rows": sheet.max_row - meaningful - (1 if headers else 0),
            "formulas": formulas,
            "merged_cells": [],
            "hidden_rows": [],
            "hidden_columns": [],
            "role": rules["role"],
            "ignored": rules["role"] == "IGNORED",
        }
    target = inventory["target questions"]
    target["separator_rows"] = mapping.data["sheets"]["target questions"]["separator_rows"]
    target["unused_rows"] = [1]
    target["structural_rows"] = 16
    for sheet_name in ("knowledge base - rag questions", "random question"):
        sheet = book[sheet_name]
        data_end = mapping.data["sheets"][sheet_name]["data_rows"][1]
        inventory[sheet_name]["formatted_blank_rows"] = sheet.max_row - data_end
    return inventory


def _domain_question_plans(target_sheet, target_rules):
    domains = []
    questions = []
    merge_rule = target_rules["merge_rules"][0]
    merge_rows = tuple(merge_rule["source_rows"])
    merge_start = merge_rows[0]
    for domain in target_rules["domains"]:
        source_rows = list(range(domain["start_row"], domain["end_row"] + 1))
        number = 0
        domain_questions = []
        index = 0
        while index < len(source_rows):
            row_number = source_rows[index]
            if row_number == merge_start:
                rows = merge_rows
                question_text = merge_rule["canonical_text"]
                index += len(merge_rows)
            else:
                rows = (row_number,)
                question_text = target_sheet.cell(row=row_number, column=1).value
                index += 1
            number += 1
            stable_id = f"{domain['code']}-{number:03d}"
            domain_questions.append(
                QuestionPlan(
                    stable_id=stable_id,
                    question_text=question_text,
                    domain_slug=domain["slug"],
                    tags=(),
                    source_role="TARGET_QUESTIONS",
                    source_rows=tuple(rows),
                    sheet_name="target questions",
                )
            )
        questions.extend(domain_questions)
        domains.append(
            DomainPlan(
                name=domain["name"],
                slug=domain["slug"],
                code=domain["code"],
                source_sheet="target questions",
                start_row=domain["start_row"],
                end_row=domain["end_row"],
                interpretation_source=target_rules["domain_interpretation_source"],
                meaningful_source_rows=len(source_rows),
                merged_source_rows=1 if merge_start in source_rows else 0,
                question_count=len(domain_questions),
                first_id=domain_questions[0].stable_id,
                last_id=domain_questions[-1].stable_id,
            )
        )
    return domains, questions


def inspect_workbook(source_path, mapping):
    path = Path(source_path)
    source_bytes = path.read_bytes()
    source_sha256 = hashlib.sha256(source_bytes).hexdigest()
    if source_sha256 != mapping.expected_sha256:
        raise ValidationError(
            "Workbook SHA-256 mismatch: "
            f"expected {mapping.expected_sha256}, found {source_sha256}."
        )

    source_book = load_workbook(path, read_only=False, data_only=False)
    displayed_book = load_workbook(path, read_only=False, data_only=True)
    try:
        _validate_workbook_structure(source_book, displayed_book, mapping)
        inventory = _build_inventory(source_book, mapping)
        source_metadata = _core_property_metadata(source_bytes)
        target_rules = mapping.data["sheets"]["target questions"]
        domains, questions = _domain_question_plans(
            source_book["target questions"], target_rules
        )
        random_rules = mapping.data["sheets"]["random question"]
        domains.append(
            DomainPlan(
                name=random_rules["domain_name"],
                slug=random_rules["domain_slug"],
                code=random_rules["domain_code"],
                source_sheet="random question",
                start_row=2,
                end_row=14,
                interpretation_source=random_rules["interpretation_source"],
                meaningful_source_rows=13,
                merged_source_rows=0,
                question_count=13,
                first_id="GEN-001",
                last_id="GEN-013",
            )
        )
        random_sheet = source_book["random question"]
        for number, row_number in enumerate(range(2, 15), start=1):
            questions.append(
                QuestionPlan(
                    stable_id=f"GEN-{number:03d}",
                    question_text=random_sheet.cell(row=row_number, column=1).value,
                    domain_slug=random_rules["domain_slug"],
                    tags=(random_rules["tag"],),
                    source_role="RANDOM_GENERALIZATION",
                    source_rows=(row_number,),
                    sheet_name="random question",
                )
            )
        kb_rules = mapping.data["sheets"]["knowledge base - rag questions"]
        kb_question = kb_rules["canonical_question"]
        domains.append(
            DomainPlan(
                name=kb_question["domain_name"],
                slug=kb_question["domain_slug"],
                code=kb_question["domain_code"],
                source_sheet="knowledge base - rag questions",
                start_row=2,
                end_row=5,
                interpretation_source=kb_question["interpretation_source"],
                meaningful_source_rows=4,
                merged_source_rows=0,
                question_count=1,
                first_id=kb_question["stable_id"],
                last_id=kb_question["stable_id"],
            )
        )
        questions.append(
            QuestionPlan(
                stable_id=kb_question["stable_id"],
                question_text=source_book["knowledge base - rag questions"].cell(
                    row=kb_question["source_row"], column=1
                ).value,
                domain_slug=kb_question["domain_slug"],
                tags=(),
                source_role="KB_RAG",
                source_rows=(2, 3, 4, 5),
                sheet_name="knowledge base - rag questions",
            )
        )

        question_by_source = {
            (question.sheet_name, row): question
            for question in questions
            for row in question.source_rows
        }
        source_rows = []
        observations = []
        warnings = []
        legacy_rules = mapping.data["legacy_observation"]

        target_source = source_book["target questions"]
        target_display = displayed_book["target questions"]
        for row_number in range(2, 203):
            question = question_by_source.get(("target questions", row_number))
            if row_number in target_rules["separator_rows"]:
                row = SourceRowPlan(
                    sheet_name="target questions",
                    row_number=row_number,
                    source_role="TARGET_QUESTIONS",
                    classification="STRUCTURAL",
                    interpretation_source="STRUCTURAL",
                    structural_kind="BLANK_SEPARATOR",
                    cells=_cell_snapshots(target_source, target_display, row_number, [""]),
                    mapping_context={"separator": True, "source_label": None},
                )
            else:
                conversation_start, conversation_end = target_rules[
                    "conversation_grouping_deferred_rows"
                ]
                row = SourceRowPlan(
                    sheet_name="target questions",
                    row_number=row_number,
                    source_role="TARGET_QUESTIONS",
                    classification="QUESTION_DEFINITION",
                    interpretation_source="PRODUCT_OWNER_MAPPING",
                    cells=_cell_snapshots(target_source, target_display, row_number, [""]),
                    mapping_context={
                        "domain_code": next(
                            domain.code for domain in domains
                            if domain.source_sheet == "target questions"
                            and domain.start_row <= row_number <= domain.end_row
                        ),
                        "canonical_source_rows": list(question.source_rows),
                        "conversation_grouping_deferred": conversation_start
                        <= row_number
                        <= conversation_end,
                    },
                    domain_slug=question.domain_slug,
                    question_stable_id=question.stable_id,
                )
            row.source_fingerprint = _row_fingerprint(row)
            source_rows.append(row)

        for sheet_name, source_role in (
            ("known questions", "KNOWN_QUESTIONS"),
            ("random question", "RANDOM_GENERALIZATION"),
            ("knowledge base - rag questions", "KB_RAG"),
        ):
            rules = mapping.data["sheets"][sheet_name]
            source_sheet = source_book[sheet_name]
            display_sheet = displayed_book[sheet_name]
            headers = rules["headers"]
            start_row, end_row = rules["data_rows"]
            for row_number in range(start_row, end_row + 1):
                values = {
                    header: source_sheet.cell(row=row_number, column=index).value
                    for index, header in enumerate(headers, start=1)
                }
                question = question_by_source.get((sheet_name, row_number))
                classification = (
                    "QUESTION_AND_OBSERVATION" if question else "LEGACY_OBSERVATION"
                )
                row = SourceRowPlan(
                    sheet_name=sheet_name,
                    row_number=row_number,
                    source_role=source_role,
                    classification=classification,
                    interpretation_source=(
                        "PRODUCT_OWNER_MAPPING" if question else "SOURCE_DATA"
                    ),
                    cells=_cell_snapshots(
                        source_sheet, display_sheet, row_number, headers
                    ),
                    mapping_context={
                        "legacy_observation": True,
                        "canonical_link_policy": (
                            "EXPLICIT_MAPPING" if question else "UNLINKED_NO_APPROVED_MAPPING"
                        ),
                    },
                    domain_slug=question.domain_slug if question else None,
                    question_stable_id=question.stable_id if question else None,
                )
                row.source_fingerprint = _row_fingerprint(row)
                source_rows.append(row)
                normalized = legacy_rules["acceptable_normalization"].get(
                    values.get("Acceptable")
                )
                experiment_metadata = {}
                observed_answer = values.get("Answer")
                comments = values.get("Comments")
                secondary_comments = None
                if sheet_name == "knowledge base - rag questions":
                    observed_answer = None
                    comments = values.get("Comment")
                    secondary_comments = values.get("Comments")
                    experiment_metadata = {
                        "answer_size": _value_as_text(values.get("Answer Size")),
                        "token": {
                            "source_value": _value_as_text(values.get("Token")),
                            "unit": None,
                            "meaning": "Unspecified beyond source column label",
                        },
                        "total_time": {
                            "source_value": _value_as_text(values.get("Total time")),
                            "unit": None,
                            "meaning": "Unspecified beyond source column label",
                        },
                    }
                observations.append(
                    LegacyObservationPlan(
                        sheet_name=sheet_name,
                        row_number=row_number,
                        question_stable_id=question.stable_id if question else None,
                        source_question_text=values["Question"],
                        observed_answer_text=observed_answer,
                        legacy_expectation=values.get("Expected Answer"),
                        acceptable_source_value=values.get("Acceptable"),
                        legacy_judgment=normalized,
                        comments=comments,
                        secondary_comments=secondary_comments,
                        experiment_metadata=experiment_metadata,
                    )
                )

        conversation_start, conversation_end = target_rules[
            "conversation_grouping_deferred_rows"
        ]
        warnings.append(
            WarningPlan(
                sheet_name="target questions",
                source_rows=tuple(range(conversation_start, conversation_end + 1)),
                code="CONVERSATION_GROUPING_DEFERRED",
                explanation=(
                    "These canonical questions are possible ordered follow-ups; M2 preserves "
                    "them individually and defers scenario/turn grouping to M8."
                ),
            )
        )
        for question in questions:
            if question.source_role != "TARGET_QUESTIONS":
                continue
            markers = [
                marker for marker in target_rules["potential_binding_markers"]
                if marker in question.question_text
            ]
            if markers:
                warnings.append(
                    WarningPlan(
                        sheet_name=question.sheet_name,
                        source_rows=question.source_rows,
                        code="POTENTIAL_BINDING_REQUIRES_REVIEW",
                        explanation=(
                            "Potential placeholder marker(s) "
                            + ", ".join(repr(marker) for marker in markers)
                            + "; no BindingDefinition was inferred."
                        ),
                    )
                )

        canonical_texts = {}
        for question in questions:
            canonical_texts.setdefault(question.question_text, []).append(question)
        known_groups = {}
        for observation in observations:
            if observation.sheet_name == "known questions":
                known_groups.setdefault(observation.source_question_text, []).append(observation)
        for text, group in known_groups.items():
            rows = tuple(observation.row_number for observation in group)
            matches = canonical_texts.get(text, [])
            if len(group) > 1 or matches:
                match_ids = [question.stable_id for question in matches]
                warnings.append(
                    WarningPlan(
                        sheet_name="known questions",
                        source_rows=rows,
                        code="AMBIGUOUS_CANONICAL_LINK",
                        explanation=(
                            "Exact repeated/source-context text was preserved without an automatic "
                            f"canonical link; candidate canonical IDs: {match_ids or 'none'}."
                        ),
                    )
                )
        for question in questions:
            if question.source_role != "RANDOM_GENERALIZATION":
                continue
            matches = [
                other.stable_id
                for other in canonical_texts.get(question.question_text, [])
                if other.stable_id != question.stable_id
            ]
            if matches:
                warnings.append(
                    WarningPlan(
                        sheet_name="random question",
                        source_rows=question.source_rows,
                        code="AMBIGUOUS_CANONICAL_LINK",
                        explanation=(
                            f"Exact text also occurs as {matches}; mapping v1 preserves the "
                            "generalization Question as a distinct identity."
                        ),
                    )
                )
        for sheet_name, row_range in (
            ("known questions", tuple(range(2, 41))),
            ("random question", tuple(range(2, 15))),
            ("knowledge base - rag questions", tuple(range(2, 6))),
        ):
            warnings.append(
                WarningPlan(
                    sheet_name=sheet_name,
                    source_rows=row_range,
                    code="LEGACY_METADATA_UNKNOWN",
                    explanation=(
                        "Historical product, environment, target, build, execution/review time, "
                        "reviewer, bindings, and session identity are absent and remain unknown."
                    ),
                )
            )
        warnings.append(
            WarningPlan(
                sheet_name="knowledge base - rag questions",
                source_rows=tuple(range(2, 6)),
                code="UNSUPPORTED_SOURCE_FIELD",
                explanation=(
                    "Token meaning and Total time unit/clock boundary are unspecified; raw values "
                    "are retained only as source experiment metadata."
                ),
            )
        )

        inventory["target questions"].update(
            {
                "candidate_canonical_questions": sum(
                    question.source_role == "TARGET_QUESTIONS"
                    for question in questions
                ),
                "candidate_legacy_observations": 0,
                "provenance_rows": sum(
                    row.sheet_name == "target questions" for row in source_rows
                ),
                "ignored_meaningful_rows": 0,
            }
        )
        inventory["random question"].update(
            {
                "candidate_canonical_questions": sum(
                    question.source_role == "RANDOM_GENERALIZATION"
                    for question in questions
                ),
                "candidate_legacy_observations": sum(
                    observation.sheet_name == "random question"
                    for observation in observations
                ),
                "provenance_rows": sum(
                    row.sheet_name == "random question" for row in source_rows
                ),
                "ignored_meaningful_rows": 0,
            }
        )
        inventory["known questions"].update(
            {
                "candidate_canonical_questions": 0,
                "candidate_legacy_observations": sum(
                    observation.sheet_name == "known questions"
                    for observation in observations
                ),
                "provenance_rows": sum(
                    row.sheet_name == "known questions" for row in source_rows
                ),
                "ignored_meaningful_rows": 0,
            }
        )
        inventory["knowledge base - rag questions"].update(
            {
                "candidate_canonical_questions": sum(
                    question.source_role == "KB_RAG" for question in questions
                ),
                "candidate_legacy_observations": sum(
                    observation.sheet_name == "knowledge base - rag questions"
                    for observation in observations
                ),
                "provenance_rows": sum(
                    row.sheet_name == "knowledge base - rag questions"
                    for row in source_rows
                ),
                "ignored_meaningful_rows": 0,
            }
        )
        inventory["Interface Issues"].update(
            {
                "candidate_canonical_questions": 0,
                "candidate_legacy_observations": 0,
                "provenance_rows": 0,
                "ignored_meaningful_rows": inventory["Interface Issues"][
                    "meaningful_rows"
                ],
                "ignored_reason": mapping.data["sheets"]["Interface Issues"][
                    "reason"
                ],
            }
        )

        return WorkbookImportPlan(
            source_path=path,
            source_filename=path.name,
            source_sha256=source_sha256,
            source_size_bytes=len(source_bytes),
            source_metadata=source_metadata,
            inventory=inventory,
            mapping=mapping,
            domains=domains,
            questions=questions,
            source_rows=source_rows,
            observations=observations,
            warnings=warnings,
            tags=(random_rules["tag"],),
        )
    finally:
        source_book.close()
        displayed_book.close()

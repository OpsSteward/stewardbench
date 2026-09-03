import hashlib

import pytest
from django.apps import apps
from django.conf import settings
from django.db import connection
from openpyxl import load_workbook

from catalog.models import (
    BindingDefinition,
    Domain,
    Environment,
    EvaluationTarget,
    HistoricalFixture,
    Product,
    Question,
    QuestionVersion,
)
from catalog.services import create_domain, create_question
from corpus.mapping import load_mapping
from corpus.models import (
    ImportedSourceRow,
    ImportWarning,
    LegacyImportBatch,
    LegacyObservation,
)
from corpus.services import CorpusImportConflict, apply_import, reconcile_import
from corpus.workbook import inspect_workbook


MAPPING_PATH = settings.BASE_DIR / "import_mappings/v2_dev_troubleshooting_v1.json"
WORKBOOK_PATH = settings.BASE_DIR / "docs/reference/v2-dev-troubleshooting.xlsx"
EXPECTED_SHA256 = "559ad19500eae888c250c38d3c7213bdf38912d611cd0dbc4dfa2556fa6ec8ff"


def database_fingerprint():
    result = {}
    with connection.cursor() as cursor:
        for table in sorted(connection.introspection.django_table_names()):
            quoted = connection.ops.quote_name(table)
            cursor.execute(
                "SELECT md5(COALESCE(string_agg(row_value, E'\\n' ORDER BY row_value), '')) "
                f"FROM (SELECT row_to_json(item)::text AS row_value FROM {quoted} item) rows"
            )
            result[table] = cursor.fetchone()[0]
        cursor.execute(
            "SELECT sequencename, last_value FROM pg_sequences "
            "WHERE schemaname = 'public' ORDER BY sequencename"
        )
        result["__sequences__"] = cursor.fetchall()
    return result


def source_value(book, sheet, row, column=1):
    return book[sheet].cell(row=row, column=column).value


@pytest.mark.acceptance
@pytest.mark.postgresql
@pytest.mark.django_db
def test_acc_import_001_real_workbook_non_destructive_reconciled_import(admin_user):
    before_sha = hashlib.sha256(WORKBOOK_PATH.read_bytes()).hexdigest()
    assert before_sha == EXPECTED_SHA256
    mapping = load_mapping(MAPPING_PATH)
    plan = inspect_workbook(WORKBOOK_PATH, mapping)

    before_dry_run = database_fingerprint()
    dry_run = reconcile_import(plan)
    after_dry_run = database_fingerprint()
    assert after_dry_run == before_dry_run
    assert dry_run["summary"]["canonical_questions"] == 199
    assert dry_run["summary"]["legacy_observations"] == 56
    assert dry_run["summary"]["target_separator_rows"] == 15
    assert dry_run["sheets"]["Interface Issues"]["ignored_meaningful_rows"] == 17

    first = apply_import(plan=plan, actor=admin_user)
    applied_fingerprint = database_fingerprint()
    second = apply_import(plan=plan, actor=admin_user)
    assert first["applied"] is True
    assert second["applied"] is False
    assert second["batch_id"] == first["batch_id"]
    assert database_fingerprint() == applied_fingerprint

    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == mapping.data["recognized_sheets"]
        target_samples = {
            2: "TRAV-001",
            8: "L0-001",
            19: "L1-001",
            27: "L2-001",
            38: "L3-001",
            50: "RUBIN-001",
            65: "DISJ-001",
            78: "CAP-001",
            89: "BLAST-001",
            106: "POWER-001",
            117: "TEMP-001",
            131: "RCA-001",
            143: "PROV-001",
            158: "DOC-001",
            168: "VIZ-001",
            184: "PROFILE-001",
            193: "PROFILE-010",
            196: "PROFILE-013",
            202: "PROFILE-019",
        }
        for row_number, stable_id in target_samples.items():
            source_row = ImportedSourceRow.objects.get(
                sheet_name="target questions", source_row_number=row_number
            )
            assert source_row.question.stable_id == stable_id
            assert source_row.question_version.question_text == source_value(
                workbook, "target questions", row_number
            )
            assert source_row.source_cells.get(
                coordinate=f"A{row_number}"
            ).raw_value == source_value(workbook, "target questions", row_number)
            assert source_row.interpretation_source == "PRODUCT_OWNER_MAPPING"

        split = ImportedSourceRow.objects.filter(
            sheet_name="target questions", source_row_number__in=(155, 156)
        ).order_by("source_row_number")
        assert list(split.values_list("question__stable_id", flat=True)) == [
            "PROV-013",
            "PROV-013",
        ]
        assert split[0].source_cells.get(coordinate="A155").raw_value == source_value(
            workbook, "target questions", 155
        )
        assert split[1].source_cells.get(coordinate="A156").raw_value == source_value(
            workbook, "target questions", 156
        )
        assert split[0].question_version.question_text == (
            "Which validation findings affect this device, interface, adjacency, "
            "service, or infrastructure segment?"
        )

        for row_number, stable_id in ((2, "GEN-001"), (4, "GEN-003"), (7, "GEN-006")):
            source_row = ImportedSourceRow.objects.get(
                sheet_name="random question", source_row_number=row_number
            )
            observation = source_row.legacy_observation
            assert source_row.question.stable_id == stable_id
            assert source_row.question_version.question_text == source_value(
                workbook, "random question", row_number, 1
            )
            assert observation.observed_answer_text == source_value(
                workbook, "random question", row_number, 2
            )
            assert observation.legacy_expectation == source_value(
                workbook, "random question", row_number, 3
            )
            assert observation.acceptable_source_value == source_value(
                workbook, "random question", row_number, 4
            )
        assert Question.objects.get(stable_id="GEN-006").current_version.question_text == (
            "Quantos transceivers no total estão em uso?"
        )
        assert Question.objects.get(stable_id="GEN-003").pk != Question.objects.get(
            stable_id="PROFILE-005"
        ).pk

        for row_number in (2, 8, 40):
            observation = LegacyObservation.objects.get(
                source_row__sheet_name="known questions",
                source_row__source_row_number=row_number,
            )
            assert observation.source_question_text == source_value(
                workbook, "known questions", row_number, 1
            )
            assert observation.observed_answer_text == source_value(
                workbook, "known questions", row_number, 2
            )
            assert observation.legacy_expectation == source_value(
                workbook, "known questions", row_number, 3
            )
            assert observation.acceptable_source_value == source_value(
                workbook, "known questions", row_number, 4
            )
            assert observation.comments == source_value(
                workbook, "known questions", row_number, 5
            )
            assert observation.question_id is None
            assert observation.question_version_id is None
        assert LegacyObservation.objects.get(
            source_row__sheet_name="known questions", source_row__source_row_number=2
        ).legacy_judgment == "GOOD"
        assert LegacyObservation.objects.get(
            source_row__sheet_name="known questions", source_row__source_row_number=8
        ).legacy_judgment == "BAD"

        for row_number in (2, 3, 4, 5):
            observation = LegacyObservation.objects.get(
                source_row__sheet_name="knowledge base - rag questions",
                source_row__source_row_number=row_number,
            )
            assert observation.question.stable_id == "KB-001"
            assert observation.observed_answer_text is None
            assert observation.experiment_metadata["answer_size"] == str(
                source_value(workbook, "knowledge base - rag questions", row_number, 2)
            )
            assert observation.comments == source_value(
                workbook, "knowledge base - rag questions", row_number, 3
            )
            assert observation.experiment_metadata["token"] == {
                "source_value": str(
                    source_value(workbook, "knowledge base - rag questions", row_number, 4)
                ),
                "unit": None,
                "meaning": "Unspecified beyond source column label",
            }
            assert observation.experiment_metadata["total_time"] == {
                "source_value": str(
                    source_value(workbook, "knowledge base - rag questions", row_number, 5)
                ),
                "unit": None,
                "meaning": "Unspecified beyond source column label",
            }
            assert observation.acceptable_source_value == source_value(
                workbook, "knowledge base - rag questions", row_number, 6
            )
            assert observation.secondary_comments == source_value(
                workbook, "knowledge base - rag questions", row_number, 7
            )
    finally:
        workbook.close()

    assert LegacyImportBatch.objects.count() == 1
    assert Domain.objects.count() == 18
    assert Question.objects.count() == QuestionVersion.objects.count() == 199
    assert Question.objects.exclude(lifecycle=Question.Lifecycle.DRAFT).count() == 0
    assert Question.objects.exclude(kind=Question.Kind.SINGLE_TURN).count() == 0
    assert LegacyObservation.objects.count() == 56
    assert BindingDefinition.objects.count() == HistoricalFixture.objects.count() == 0
    assert Product.objects.count() == Environment.objects.count() == 0
    assert EvaluationTarget.objects.count() == 0
    assert not ImportedSourceRow.objects.filter(sheet_name="Interface Issues").exists()
    assert ImportedSourceRow.objects.filter(
        sheet_name="target questions", structural_kind="BLANK_SEPARATOR"
    ).count() == 15
    conversation_warning = ImportWarning.objects.get(code="CONVERSATION_GROUPING_DEFERRED")
    assert conversation_warning.source_rows == list(range(193, 203))
    assert ImportedSourceRow.objects.filter(
        sheet_name="target questions",
        source_row_number__range=(193, 202),
        mapping_context__conversation_grouping_deferred=True,
    ).count() == 10
    assert not apps.is_installed("executions")
    assert QuestionVersion.objects.exclude(evaluation_guidance="").count() == 0
    assert all(
        "reviewer" in values
        for values in LegacyObservation.objects.values_list("unknown_metadata", flat=True)
    )
    assert hashlib.sha256(WORKBOOK_PATH.read_bytes()).hexdigest() == before_sha


@pytest.mark.acceptance
@pytest.mark.postgresql
@pytest.mark.django_db
def test_acc_import_001_manual_identity_conflict_cannot_overwrite(admin_user):
    mapping = load_mapping(MAPPING_PATH)
    plan = inspect_workbook(WORKBOOK_PATH, mapping)
    domain = create_domain(
        actor=admin_user, slug="service-traversal", name="Service Traversal"
    )
    manual = create_question(
        actor=admin_user,
        stable_id="TRAV-001",
        kind=Question.Kind.SINGLE_TURN,
        lifecycle=Question.Lifecycle.DRAFT,
        domain=domain,
        tags=(),
        rationale="Manual catalog evidence",
        question_text="A deliberately different manual definition",
    )

    report = reconcile_import(plan)
    assert report["fatal"] is True
    assert report["conflicts"] == [
        {
            "code": "SOURCE_CONFLICT",
            "object": "Question TRAV-001",
            "explanation": (
                "Existing Question does not exactly match the approved imported kind, "
                "current text, domain, and required deterministic tags; it was not "
                "changed."
            ),
        }
    ]
    with pytest.raises(CorpusImportConflict):
        apply_import(plan=plan, actor=admin_user)

    manual.refresh_from_db()
    assert manual.current_version.question_text == "A deliberately different manual definition"
    assert LegacyImportBatch.objects.count() == 0
    assert Question.objects.count() == 1
